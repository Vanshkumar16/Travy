"""Render data/jd/samples/*.txt as PDF and DOCX into data/jd/fixtures/ for upload testing."""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "data" / "jd" / "samples"
OUT = ROOT / "data" / "jd" / "fixtures"


def blocks(text):
    """Yield ("heading"|"bullets"|"para", payload) from a sample JD."""
    bullets = []
    for line in text.splitlines():
        line = line.strip()
        if line.startswith("•"):
            bullets.append(line.lstrip("• ").strip())
            continue
        if bullets:
            yield "bullets", bullets
            bullets = []
        if not line:
            continue
        is_heading = len(line.split()) <= 6 and not line.endswith(".") and ":" not in line
        yield ("heading" if is_heading else "para"), line
    if bullets:
        yield "bullets", bullets


def to_pdf(text, path):
    styles = getSampleStyleSheet()
    story = []
    for kind, payload in blocks(text):
        if kind == "heading":
            story += [Spacer(1, 6), Paragraph(payload, styles["Heading3"])]
        elif kind == "para":
            story.append(Paragraph(payload, styles["BodyText"]))
        else:
            story.append(ListFlowable([ListItem(Paragraph(b, styles["BodyText"])) for b in payload], bulletType="bullet"))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def to_docx(text, path):
    doc = Document()
    for kind, payload in blocks(text):
        if kind == "heading":
            doc.add_heading(payload, level=2)
        elif kind == "para":
            doc.add_paragraph(payload)
        else:
            for b in payload:
                doc.add_paragraph(b, style="List Bullet")
    doc.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for txt in sorted(SAMPLES.glob("*.txt")):
        text = txt.read_text(encoding="utf-8")
        to_pdf(text, OUT / f"{txt.stem}.pdf")
        to_docx(text, OUT / f"{txt.stem}.docx")
        print("wrote", txt.stem)
