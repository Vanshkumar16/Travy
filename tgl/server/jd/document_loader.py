"""Turn an uploaded JD (PDF / DOCX / TXT) into clean plain text."""

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedDocumentError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedDocumentError(
            f"Unsupported file type '{ext or filename}'. Upload a PDF, DOCX or TXT file."
        )
    text = normalize_text(text)
    if len(text) < 40:
        raise UnsupportedDocumentError(
            "Could not read enough text from this file. If it is a scanned PDF, export it as text first."
        )
    return text


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # Word list paragraphs lose their bullet glyph; restore it so the
        # section parser sees them as list items.
        style = (para.style.name or "").lower() if para.style is not None else ""
        if "list" in style and not text.startswith(("•", "-", "*")):
            text = f"• {text}"
        lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


_BULLET_RE = re.compile(r"^\s*(?:[•●▪◦‣∙·\-\*–]|\d+[.)])\s+", re.MULTILINE)


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = text.replace("’", "'").replace("“", '"').replace("”", '"')
    text = _BULLET_RE.sub("• ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
